"""Exporta el acta a Word (.docx) EDITABLE para que la secretaria remate la redacción.

Convierte `data/actas/<id>.acta.txt` (el acta ya estructurada) en un .docx con título, escudo del
municipio en el membrete, secciones (Asistentes, Orden del día, Estadísticas) y el cuerpo con cada
turno como "Orador:" en negrita + su texto. Todo editable en Word, que es donde trabaja la secretaria.

Uso:  python -m src.acta_docx <video_id>   ->  data/actas/<id>.acta.docx
"""
from __future__ import annotations

import re

from .config import resolve_path

_TURN = re.compile(r"^\[(\d+:\d\d:\d\d)\]\s+(.+?):\s*$")   # "[0:00:38] Voz 1:"


def build_docx(video_id: str) -> str:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    txt_path = resolve_path(f"data/actas/{video_id}.acta.txt")
    if not txt_path.exists():
        raise FileNotFoundError(txt_path)
    lines = txt_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    from src.acta import _find_escudo
    esc = _find_escudo(lines[1] if len(lines) > 1 else "")
    if esc:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(esc, width=Cm(2.2))
        except Exception:
            pass

    # --- cabecera: las primeras líneas hasta el primer separador o turno ---
    i = 0
    title_done = False
    while i < len(lines):
        ln = lines[i].rstrip()
        if set(ln) == {"="} and ln:                       # separador de sección
            break
        if _TURN.match(ln):
            break
        if ln.strip():
            if not title_done:
                h = doc.add_heading(ln.strip(), level=0); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_done = True
            else:
                p = doc.add_paragraph(ln.strip()); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1

    # --- resto: secciones (====) y cuerpo (turnos) ---
    cur = None                                             # párrafo del turno en curso
    while i < len(lines):
        ln = lines[i].rstrip()
        if set(ln) == {"="} and ln:                        # delimitador -> la línea siguiente es título
            if i + 1 < len(lines) and lines[i + 1].strip():
                doc.add_heading(lines[i + 1].strip().title(), level=1)
                i += 2
                while i < len(lines) and set(lines[i].rstrip()) == {"="}:
                    i += 1
                cur = None
                continue
            i += 1
            continue
        m = _TURN.match(ln)
        if m:                                              # nuevo turno: "Orador:" en negrita
            p = doc.add_paragraph()
            r = p.add_run(f"{m.group(2)}  ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x7A, 0x1F, 0x2B)
            small = p.add_run(f"[{m.group(1)}]")
            small.italic = True; small.font.size = Pt(8); small.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            cur = doc.add_paragraph(); cur.paragraph_format.left_indent = Cm(0.5)
        elif ln.strip():
            if cur is not None:                            # cuerpo del turno (continúa)
                cur.add_run((" " if cur.runs else "") + ln.strip())
            else:
                doc.add_paragraph(ln.strip())
        i += 1

    out = resolve_path(f"data/actas/{video_id}.acta.docx")
    doc.save(str(out))
    return str(out)


def main() -> None:
    import sys
    vid = sys.argv[1] if len(sys.argv) > 1 else "madrid2432"
    print("->", build_docx(vid))


if __name__ == "__main__":
    main()
